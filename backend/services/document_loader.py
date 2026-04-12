"""文档解析服务 - 支持 PDF/TXT/MD/Word/Excel"""
from pathlib import Path
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from config import CHUNK_SIZE, CHUNK_OVERLAP

text_splitter = RecursiveCharacterTextSplitter(
    chunk_size=CHUNK_SIZE,
    chunk_overlap=CHUNK_OVERLAP,
)


def load_txt(file_path: str) -> str:
    return Path(file_path).read_text(encoding="utf-8")


def load_markdown(file_path: str) -> str:
    return Path(file_path).read_text(encoding="utf-8")


def load_pdf(file_path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(file_path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def load_docx(file_path: str) -> str:
    from docx import Document as DocxDocument
    doc = DocxDocument(file_path)
    return "\n".join(p.text for p in doc.paragraphs)


def load_excel(file_path: str) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(file_path, read_only=True)
    lines = []
    for sheet in wb.sheetnames:
        ws = wb[sheet]
        for row in ws.iter_rows(values_only=True):
            lines.append("\t".join(str(c) if c is not None else "" for c in row))
    wb.close()
    return "\n".join(lines)


LOADERS = {
    ".txt": load_txt,
    ".md": load_markdown,
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".doc": load_docx,
    ".xlsx": load_excel,
    ".xls": load_excel,
}

SUPPORTED_EXTENSIONS = set(LOADERS.keys())


def parse_and_split(file_path: str, knowledge_base_id: str) -> list[Document]:
    """解析文件并切分为文档块"""
    ext = Path(file_path).suffix.lower()
    loader = LOADERS.get(ext)
    if not loader:
        raise ValueError(f"不支持的文件格式: {ext}")

    text = loader(file_path)
    if not text.strip():
        return []

    chunks = text_splitter.split_text(text)
    # 过滤空字符串和纯空白分块，避免 Embedding API 400 错误
    valid_chunks = [c for c in chunks if c.strip()]
    return [
        Document(
            page_content=chunk,
            metadata={
                "source": Path(file_path).name,
                "knowledge_base_id": knowledge_base_id,
                "chunk_index": i,
            },
        )
        for i, chunk in enumerate(valid_chunks)
    ]
